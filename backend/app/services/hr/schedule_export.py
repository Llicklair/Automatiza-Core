"""Export de horarios de trabajo a Excel (openpyxl) y PDF (reportlab)."""

import io
from datetime import date

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models.hr import WorkSchedule
from app.db.models.models import Employee

DAY_NAMES = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]


async def fetch_schedule_grid(
    db: AsyncSession, tenant_id, employee_id=None
) -> list[dict]:
    """Devuelve [{employee: str, days: {0..6: 'HH:MM-HH:MM' | None}}] ordenado por nombre."""
    emp_q = select(Employee).where(Employee.tenant_id == tenant_id)
    if employee_id is not None:
        emp_q = emp_q.where(Employee.id == employee_id)
    employees = (await db.execute(emp_q.order_by(Employee.name))).scalars().all()

    sched_q = select(WorkSchedule).where(
        WorkSchedule.tenant_id == tenant_id, WorkSchedule.active.is_(True)
    )
    if employee_id is not None:
        sched_q = sched_q.where(WorkSchedule.employee_id == employee_id)
    schedules = (await db.execute(sched_q)).scalars().all()

    by_emp: dict = {}
    for s in schedules:
        by_emp.setdefault(s.employee_id, {})[s.day_of_week] = f"{s.start_time}-{s.end_time}"

    return [
        {"employee": e.name, "days": by_emp.get(e.id, {})}
        for e in employees
        if employee_id is not None or by_emp.get(e.id)
    ]


def build_schedules_xlsx(grid: list[dict]) -> bytes:
    """Excel: una fila por empleado, columnas Lunes..Domingo."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Horarios"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Border(*([Side(style="thin", color="B0B0B0")] * 4))

    headers = ["Empleado", *DAY_NAMES]
    ws.append(headers)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
        cell.alignment = Alignment(horizontal="center")

    for row in grid:
        ws.append([row["employee"], *[row["days"].get(d, "—") for d in range(7)]])
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.border = thin
            cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 28
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_schedules_pdf(grid: list[dict], company_name: str = "") -> bytes:
    """PDF apaisado-friendly con tabla de horarios semanales."""
    from app.services.documents._pdf_base import (
        REPORTLAB_AVAILABLE,
        _trad_table_style,
        _traditional_styles,
    )

    if not REPORTLAB_AVAILABLE:
        lines = ["HORARIOS DE TRABAJO", ""]
        for row in grid:
            days = ", ".join(
                f"{DAY_NAMES[d]} {row['days'][d]}" for d in sorted(row["days"])
            )
            lines.append(f"{row['employee']}: {days or 'sin horario'}")
        return "\n".join(lines).encode("utf-8")

    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table

    S = _traditional_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )
    title = "Horarios de trabajo" + (f" — {company_name}" if company_name else "")
    elements = [
        Paragraph(title, S["title"]),
        Paragraph(f"Generado el {date.today().strftime('%d/%m/%Y')}", S["small"]),
        Spacer(1, 6 * mm),
    ]
    data = [["Empleado", *DAY_NAMES]]
    for row in grid:
        data.append([row["employee"], *[row["days"].get(d, "—") for d in range(7)]])
    table = Table(data, colWidths=[60 * mm, *([28 * mm] * 7)], repeatRows=1)
    table.setStyle(_trad_table_style())
    elements.append(table)
    doc.build(elements)
    return buf.getvalue()
