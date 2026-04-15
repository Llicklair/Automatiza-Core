"""
Herramientas del agente de Excel.

DB fetchers, utilidades de escritura/lectura, y las tools LangChain
que el agente invoca de forma autónoma.
"""

import json
import logging
import os
import uuid
from datetime import datetime, timezone

import openpyxl
import pandas as pd
from langchain_core.tools import tool
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.future import select

from app.agents.agent_tools.documents import (
    create_document,
    get_document_content,
    list_tenant_documents,
)
from app.agents.agent_tools.knowledge import get_tenant_knowledge, upsert_tenant_knowledge
from app.db.base import AsyncSessionLocal
from app.db.models.accounting import BankTransaction
from app.db.models.billing import Invoice
from app.db.models.crm import Client
from app.db.models.hr import Employee, Payroll
from app.db.models.inventory import Product
from app.db.models.tenant import TenantDocument

logger = logging.getLogger(__name__)

UPLOADS_DIR = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "..", "..", "..", "uploads")
)
os.makedirs(UPLOADS_DIR, exist_ok=True)


# ─── DB Fetchers (reutilizados por las tools) ────────────────────────────────


async def _fetch_invoices(tenant_id: str) -> pd.DataFrame:
    async with AsyncSessionLocal() as db:
        result = await db.execute(
            select(Invoice, Client.name.label("client_name"))
            .join(Client, Invoice.client_id == Client.id)
            .where(Invoice.tenant_id == uuid.UUID(tenant_id))
            .order_by(Invoice.date.desc())
            .limit(500)
        )
        rows = result.all()
    if not rows:
        return pd.DataFrame(
            columns=["Número", "Cliente", "Fecha", "Base (€)", "IVA (€)", "Total (€)", "Estado"]
        )
    return pd.DataFrame(
        [
            {
                "Número": inv.invoice_number or "",
                "Cliente": cn or "",
                "Fecha": inv.date.strftime("%d/%m/%Y") if inv.date else "",
                "Base (€)": float(inv.amount_base or 0),
                "IVA (€)": float(inv.tax_amount or 0),
                "Total (€)": float(inv.amount_total or 0),
                "Estado": inv.status or "",
            }
            for inv, cn in rows
        ]
    )


async def _fetch_clients(tenant_id: str) -> pd.DataFrame:
    async with AsyncSessionLocal() as db:
        result = await db.execute(
            select(Client)
            .where(Client.tenant_id == uuid.UUID(tenant_id))
            .order_by(Client.name)
            .limit(500)
        )
        clients = result.scalars().all()
    if not clients:
        return pd.DataFrame(columns=["Nombre", "NIF/CIF", "Email", "Tipo"])
    return pd.DataFrame(
        [
            {
                "Nombre": c.name or "",
                "NIF/CIF": c.nif or "",
                "Email": c.email or "",
                "Tipo": c.client_type or "",
            }
            for c in clients
        ]
    )


async def _fetch_employees(tenant_id: str) -> pd.DataFrame:
    async with AsyncSessionLocal() as db:
        result = await db.execute(
            select(Employee)
            .where(Employee.tenant_id == uuid.UUID(tenant_id))
            .order_by(Employee.name)
        )
        employees = result.scalars().all()
    if not employees:
        return pd.DataFrame(
            columns=[
                "Nombre",
                "NIF",
                "Departamento",
                "Puesto",
                "Salario Base (€)",
                "IRPF (%)",
                "Estado",
            ]
        )
    return pd.DataFrame(
        [
            {
                "Nombre": e.name or "",
                "NIF": e.nif or "",
                "Departamento": e.department or "",
                "Puesto": e.role or "",
                "Salario Base (€)": float(e.base_salary or 0),
                "IRPF (%)": float(e.irpf_rate or 0),
                "Estado": e.status or "",
            }
            for e in employees
        ]
    )


async def _fetch_payrolls(tenant_id: str) -> pd.DataFrame:
    async with AsyncSessionLocal() as db:
        result = await db.execute(
            select(Payroll, Employee.name.label("emp_name"))
            .join(Employee, Payroll.employee_id == Employee.id)
            .where(Payroll.tenant_id == uuid.UUID(tenant_id))
            .order_by(Payroll.issue_date.desc())
            .limit(500)
        )
        rows = result.all()
    if not rows:
        return pd.DataFrame(
            columns=[
                "Empleado",
                "Período",
                "Salario Base (€)",
                "SS (€)",
                "IRPF (€)",
                "Neto (€)",
                "Estado",
            ]
        )
    return pd.DataFrame(
        [
            {
                "Empleado": en or "",
                "Período": f"{p.period_start.strftime('%d/%m/%Y')} - {p.period_end.strftime('%d/%m/%Y')}"
                if p.period_start
                else "",
                "Salario Base (€)": float(p.base_salary or 0),
                "SS (€)": float(
                    (p.ss_contingencias_comunes or 0)
                    + (p.ss_desempleo or 0)
                    + (p.ss_formacion_profesional or 0)
                    + (p.ss_mei or 0)
                ),
                "IRPF (€)": float(p.irpf or 0),
                "Neto (€)": float(p.net_salary or 0),
                "Estado": p.status or "",
            }
            for p, en in rows
        ]
    )


async def _fetch_products(tenant_id: str) -> pd.DataFrame:
    async with AsyncSessionLocal() as db:
        result = await db.execute(
            select(Product)
            .where(Product.tenant_id == uuid.UUID(tenant_id))
            .order_by(Product.name)
            .limit(500)
        )
        products = result.scalars().all()
    if not products:
        return pd.DataFrame(columns=["Nombre", "SKU", "Tipo", "Precio (€)", "Stock"])
    return pd.DataFrame(
        [
            {
                "Nombre": p.name or "",
                "SKU": p.sku or "",
                "Tipo": p.item_type or "",
                "Precio (€)": float(p.price or 0),
                "Stock": int(p.stock_quantity or 0),
            }
            for p in products
        ]
    )


async def _fetch_bank(tenant_id: str) -> pd.DataFrame:
    async with AsyncSessionLocal() as db:
        result = await db.execute(
            select(BankTransaction)
            .where(BankTransaction.tenant_id == uuid.UUID(tenant_id))
            .order_by(BankTransaction.date.desc())
            .limit(500)
        )
        txs = result.scalars().all()
    if not txs:
        return pd.DataFrame(columns=["Fecha", "Concepto", "Importe (€)", "Saldo (€)", "Estado"])
    return pd.DataFrame(
        [
            {
                "Fecha": t.date.strftime("%d/%m/%Y") if t.date else "",
                "Concepto": t.description or "",
                "Importe (€)": float(t.amount or 0),
                "Saldo (€)": float(t.balance or 0),
                "Estado": t.status or "",
            }
            for t in txs
        ]
    )


_FETCHER_MAP = {
    "facturas": _fetch_invoices,
    "clientes": _fetch_clients,
    "empleados": _fetch_employees,
    "nominas": _fetch_payrolls,
    "productos": _fetch_products,
    "banco": _fetch_bank,
}

_INTENT_MAP = [
    (["factura", "invoice", "venta", "cobro", "ingreso"], "facturas"),
    (["empleado", "trabajador", "plantilla", "rrhh", "personal"], "empleados"),
    (["nómina", "nomina", "salario", "sueldo"], "nominas"),
    (["cliente", "client", "cuenta", "crm"], "clientes"),
    (["producto", "catalogo", "catálogo", "inventario", "stock"], "productos"),
    (["banco", "bank", "movimiento", "transaccion"], "banco"),
]


def _detect_datasets(intent: str) -> list[str]:
    intent_lower = intent.lower()
    matched = []
    for keywords, key in _INTENT_MAP:
        if any(kw in intent_lower for kw in keywords):
            if key not in matched:
                matched.append(key)
    return matched or ["facturas"]


def _hex_to_lighter(hex_color: str, factor: float = 0.4) -> str:
    """Return a lighter version of a hex color by blending toward white."""
    h = hex_color.lstrip("#")
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    r = int(r + (255 - r) * factor)
    g = int(g + (255 - g) * factor)
    b = int(b + (255 - b) * factor)
    return f"{r:02X}{g:02X}{b:02X}"


def _write_excel(
    sheets: dict[str, pd.DataFrame], output_path: str, theme: dict | None = None
) -> None:
    _theme = theme or {}
    accent_hex = _theme.get("accent_color", "#1F4E79").lstrip("#")
    table_style = _theme.get("table_style", "striped")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(fill_type="solid", fgColor=accent_hex)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_side = Side(style="thin", color="D9D9D9")
    cell_border = Border(left=thin_side, right=thin_side, bottom=thin_side)

    if table_style == "minimal":
        alt_fill = None
    elif table_style == "bold":
        alt_fill = PatternFill(fill_type="solid", fgColor=_hex_to_lighter(accent_hex))
    else:  # "striped" (default)
        alt_fill = PatternFill(fill_type="solid", fgColor="EBF3FB")

    for sheet_name, df in sheets.items():
        ws = wb.create_sheet(title=sheet_name[:31])
        if df.empty:
            ws.append(["Sin datos"])
            continue
        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = cell_border
        for row_idx, row in enumerate(df.itertuples(index=False), start=2):
            fill = alt_fill if (alt_fill and row_idx % 2 == 0) else None
            for col_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = cell_border
                cell.alignment = Alignment(horizontal="left", vertical="center")
                if fill:
                    cell.fill = fill
        for col_idx, col_name in enumerate(df.columns, start=1):
            max_len = max(
                len(str(col_name)),
                df.iloc[:, col_idx - 1].astype(str).str.len().max() if not df.empty else 0,
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 50)
        ws.freeze_panes = "A2"

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)


# ─── Mapeo de columnas para importación ──────────────────────────────────────

_IMPORT_COLUMN_MAP = {
    "clientes": {
        "model": "Client",
        "fields": {
            "nombre": "name",
            "name": "name",
            "nif": "nif",
            "cif": "nif",
            "email": "email",
            "correo": "email",
            "telefono": "phone",
            "teléfono": "phone",
            "phone": "phone",
            "direccion": "address",
            "dirección": "address",
            "address": "address",
        },
        "required": ["name"],
    },
    "productos": {
        "model": "Product",
        "fields": {
            "nombre": "name",
            "name": "name",
            "descripcion": "description",
            "descripción": "description",
            "description": "description",
            "precio": "price",
            "price": "price",
            "pvp": "price",
            "tipo": "item_type",
            "type": "item_type",
            "iva": "tax_percentage",
            "tax": "tax_percentage",
        },
        "required": ["name"],
    },
    "empleados": {
        "model": "Employee",
        "fields": {
            "nombre": "name",
            "name": "name",
            "nif": "nif",
            "dni": "nif",
            "email": "email",
            "correo": "email",
            "cargo": "role",
            "puesto": "role",
            "role": "role",
            "departamento": "department",
            "department": "department",
            "salario": "base_salary",
            "salario_base": "base_salary",
            "base_salary": "base_salary",
        },
        "required": ["name"],
    },
}


# ─── Herramientas del agente ──────────────────────────────────────────────────


@tool
async def export_erp_data(tenant_id: str, datasets: str = "todos", user_request: str = "") -> str:
    """
    Exporta datos del ERP a un archivo Excel formateado (.xlsx).
    Genera un archivo con hojas separadas por cada tipo de dato solicitado.

    Args:
        tenant_id: ID del tenant
        datasets: Tipos de datos a exportar, separados por coma.
                  Opciones: facturas, clientes, empleados, nominas, productos, banco.
                  Ejemplo: "facturas,nominas" o "todos" para exportar todo.
        user_request: El mensaje original del usuario (para detectar datasets que falten).
    """
    return await _export_erp_data_async(tenant_id, datasets, user_request)


async def _export_erp_data_async(tenant_id: str, datasets_str: str, user_request: str = "") -> str:
    os.makedirs(UPLOADS_DIR, exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    output_filename = f"informe_{ts}.xlsx"
    output_path = os.path.join(UPLOADS_DIR, output_filename)

    if datasets_str.strip().lower() == "todos":
        dataset_keys = list(_FETCHER_MAP.keys())
    else:
        # Primero intentar claves exactas, luego detección por intención
        raw_keys = [d.strip().lower() for d in datasets_str.split(",")]
        dataset_keys = []
        for key in raw_keys:
            if key in _FETCHER_MAP:
                dataset_keys.append(key)
            else:
                # Intentar detectar por palabras clave
                detected = _detect_datasets(key)
                for dk in detected:
                    if dk not in dataset_keys:
                        dataset_keys.append(dk)

        # Salvaguardia: detectar datasets del prompt original que el LLM haya omitido
        if user_request:
            from_intent = _detect_datasets(user_request)
            for dk in from_intent:
                if dk not in dataset_keys:
                    logger.info("Dataset '%s' detectado del prompt original, añadiendo", dk)
                    dataset_keys.append(dk)

    sheets: dict[str, pd.DataFrame] = {}
    for key in dataset_keys:
        fetcher = _FETCHER_MAP.get(key)
        if fetcher:
            df = await fetcher(tenant_id)
            sheets[key.capitalize()] = df

    if not sheets:
        return f"Error: No se reconocieron los datasets '{datasets_str}'. Opciones: {', '.join(_FETCHER_MAP.keys())}"

    from app.services.template_service import get_default_theme

    _theme = None
    try:
        async with AsyncSessionLocal() as _db:
            _theme = await get_default_theme(uuid.UUID(tenant_id), "excel", _db)
    except Exception as _e:
        logger.warning("Error cargando tema Excel para tenant %s: %s", tenant_id, _e)

    _write_excel(sheets, output_path, theme=_theme)

    # Registrar en BD
    async with AsyncSessionLocal() as db:
        doc = TenantDocument(
            tenant_id=uuid.UUID(tenant_id),
            file_name=output_filename,
            file_path=output_path,
            file_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            file_size=os.path.getsize(output_path),
            category="excels",
            status="completed",
            parsed_content=f"Excel exportado: {', '.join(sheets.keys())}",
        )
        db.add(doc)
        await db.commit()

    total_rows = sum(len(df) for df in sheets.values())
    sheet_summary = ", ".join(f"{name} ({len(df)} registros)" for name, df in sheets.items())
    return (
        f"Excel generado correctamente.\n"
        f"Archivo: {output_filename}\n"
        f"Hojas: {sheet_summary}\n"
        f"Total: {total_rows} registros\n"
        f"El archivo está disponible en la sección de Documentos."
    )


@tool
async def list_available_datasets(tenant_id: str) -> str:
    """
    Lista los tipos de datos disponibles para exportar a Excel y cuántos registros tiene cada uno.
    Útil para que el usuario sepa qué puede exportar antes de pedirlo.

    Args:
        tenant_id: ID del tenant
    """
    return await _list_available_datasets_async(tenant_id)


async def _list_available_datasets_async(tenant_id: str) -> str:
    lines = []
    for key, fetcher in _FETCHER_MAP.items():
        try:
            df = await fetcher(tenant_id)
            lines.append(f"- {key}: {len(df)} registros")
        except Exception:
            logger.debug("Error consultando dataset %s", key, exc_info=True)
            lines.append(f"- {key}: error al consultar")
    return "Datasets disponibles para exportar:\n" + "\n".join(lines)


async def _load_excel_doc(
    tenant_id: str, document_id: str
) -> "tuple[TenantDocument | None, str | None]":
    """Carga TenantDocument desde BD y valida que el archivo exista en disco."""
    async with AsyncSessionLocal() as db:
        result = await db.execute(
            select(TenantDocument).where(
                TenantDocument.tenant_id == uuid.UUID(tenant_id),
                TenantDocument.id == uuid.UUID(document_id),
            )
        )
        doc = result.scalar_one_or_none()
    if not doc:
        return None, f"Error: Documento {document_id} no encontrado."
    if not doc.file_path or not os.path.exists(doc.file_path):
        return None, f"Error: Archivo no encontrado en disco: {doc.file_path}"
    return doc, None


def _detect_import_target(sheet_title: str) -> str | None:
    """Infiere el target de importacion a partir del nombre de la hoja."""
    title = sheet_title.lower().strip()
    for key in _IMPORT_COLUMN_MAP:
        if key in title:
            return key
    return None


def _map_excel_columns(headers: list[str], config: dict) -> dict[int, str]:
    """Mapea indices de columna Excel a nombres de campo del modelo."""
    return {idx: config["fields"][h] for idx, h in enumerate(headers) if h in config["fields"]}


async def _import_rows_to_db(
    rows: list, col_mapping: dict, config: dict, model_cls, tenant_id: str
) -> tuple[int, int, list[str]]:
    """Persiste filas Excel en la BD. Devuelve (created, skipped, errors)."""
    created, skipped, errors = 0, 0, []
    async with AsyncSessionLocal() as db:
        for row_idx, row in enumerate(rows[1:], start=2):
            record_data: dict = {"tenant_id": uuid.UUID(tenant_id)}
            for col_idx, field_name in col_mapping.items():
                value = row[col_idx] if col_idx < len(row) else None
                if value is not None:
                    record_data[field_name] = value
            if any(not record_data.get(r) for r in config["required"]):
                skipped += 1
                continue
            try:
                db.add(model_cls(**record_data))
                created += 1
            except Exception as e:
                errors.append(f"Fila {row_idx}: {e}")
                skipped += 1
        if created > 0:
            await db.commit()
    return created, skipped, errors


def _parse_mods_json(modifications_json: str) -> list | str:
    """Parsea el JSON de modificaciones. Devuelve la lista o un mensaje de error."""
    try:
        mods = json.loads(modifications_json)
    except (json.JSONDecodeError, TypeError):
        return 'Error: El formato de modificaciones no es JSON válido. Ejemplo: [{"cell": "B3", "value": 1500}]'
    if not isinstance(mods, list):
        return "Error: Las modificaciones deben ser una lista JSON."
    return mods


def _apply_mods_to_workbook(wb, mods: list, default_sheet: str) -> tuple[int, list[str]]:
    """Aplica una lista de modificaciones a un workbook openpyxl. Devuelve (applied, errors)."""
    applied, errors = 0, []
    for mod in mods:
        cell_ref = mod.get("cell", "")
        value = mod.get("value")
        sheet = mod.get("sheet", default_sheet or wb.sheetnames[0])
        if not cell_ref:
            errors.append("Modificación sin 'cell' especificada.")
            continue
        if sheet not in wb.sheetnames:
            errors.append(f"Hoja '{sheet}' no existe. Disponibles: {', '.join(wb.sheetnames)}")
            continue
        try:
            wb[sheet][cell_ref] = value
            applied += 1
        except Exception as e:
            errors.append(f"Error en {sheet}!{cell_ref}: {e}")
    return applied, errors


@tool
async def import_excel(
    tenant_id: str, document_id: str, target: str = "", sheet_name: str = ""
) -> str:
    """
    Importa datos desde un archivo Excel (.xlsx) subido al sistema hacia la base de datos del ERP.
    Lee las columnas del Excel y las mapea automáticamente a campos de la BD.

    Args:
        tenant_id: ID del tenant
        document_id: ID del documento Excel a importar (usar list_tenant_documents para obtenerlo)
        target: Tipo de datos a importar ('clientes', 'productos', 'empleados'). Si vacío, se detecta por nombre de hoja.
        sheet_name: Nombre de la hoja a importar (vacío = primera hoja)
    """
    return await _import_excel_async(tenant_id, document_id, target, sheet_name)


async def _import_excel_async(
    tenant_id: str, document_id: str, target: str, sheet_name: str
) -> str:
    MODEL_MAP = {"Client": Client, "Product": Product, "Employee": Employee}

    try:
        doc, err = await _load_excel_doc(tenant_id, document_id)
        if err:
            return err

        wb = openpyxl.load_workbook(doc.file_path, read_only=True, data_only=True)

        # Seleccionar hoja
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                return f"Error: Hoja '{sheet_name}' no encontrada. Hojas disponibles: {', '.join(wb.sheetnames)}"
            ws = wb[sheet_name]
        else:
            ws = wb.active

        # Detectar target
        target = (target or "").strip().lower() or _detect_import_target(ws.title) or ""
        if not target:
            return (
                f"Error: No se pudo detectar el tipo de datos de la hoja '{ws.title}'. "
                f"Especifica target: 'clientes', 'productos' o 'empleados'."
            )
        if target not in _IMPORT_COLUMN_MAP:
            return f"Error: Target '{target}' no soportado. Opciones: {', '.join(_IMPORT_COLUMN_MAP.keys())}"

        config = _IMPORT_COLUMN_MAP[target]
        model_cls = MODEL_MAP[config["model"]]

        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if len(rows) < 2:
            return "Error: El archivo no tiene datos (solo headers o vacío)."

        headers = [str(h).strip().lower() if h else "" for h in rows[0]]
        col_mapping = _map_excel_columns(headers, config)
        if not col_mapping:
            return (
                f"Error: No se reconocieron columnas para '{target}'. "
                f"Columnas encontradas: {', '.join(h for h in headers if h)}. "
                f"Columnas esperadas: {', '.join(config['fields'].keys())}"
            )

        sheet_title = ws.title
        created, skipped, errors = await _import_rows_to_db(rows, col_mapping, config, model_cls, tenant_id)

        result_lines = [
            f"Importación completada desde '{sheet_title}'.",
            f"Target: {target}",
            f"Creados: {created} registros",
            f"Omitidos: {skipped} filas (datos incompletos o errores)",
        ]
        if errors[:5]:
            result_lines.append(f"Errores: {'; '.join(errors[:5])}")
        return "\n".join(result_lines)
    except Exception as e:
        return f"Error importando Excel: {e}"


@tool
async def modify_excel(
    tenant_id: str,
    document_id: str,
    modifications: str = "",
    sheet_name: str = "",
) -> str:
    """
    Modifica celdas de un archivo Excel existente. Acepta instrucciones en formato JSON.
    Cada modificación especifica hoja (opcional), celda y nuevo valor.

    Args:
        tenant_id: ID del tenant
        document_id: ID del documento Excel a modificar (usar list_tenant_documents para obtenerlo)
        modifications: JSON con las modificaciones. Formato:
            [{"cell": "B3", "value": 1500}, {"cell": "A1", "value": "Nuevo título"}]
            o con hoja: [{"sheet": "Facturas", "cell": "C5", "value": 2000}]
        sheet_name: Hoja por defecto donde aplicar modificaciones (vacío = primera hoja)
    """
    return await _modify_excel_async(tenant_id, document_id, modifications, sheet_name)


async def _modify_excel_async(
    tenant_id: str,
    document_id: str,
    modifications_json: str,
    default_sheet: str,
) -> str:
    try:
        doc, err = await _load_excel_doc(tenant_id, document_id)
        if err:
            return err

        mods = _parse_mods_json(modifications_json)
        if isinstance(mods, str):
            return mods

        wb = openpyxl.load_workbook(doc.file_path)
        applied, errors = _apply_mods_to_workbook(wb, mods, default_sheet)
        wb.save(doc.file_path)
        wb.close()

        async with AsyncSessionLocal() as db:
            result = await db.execute(
                select(TenantDocument).where(TenantDocument.id == uuid.UUID(document_id))
            )
            doc_upd = result.scalar_one_or_none()
            if doc_upd:
                doc_upd.file_size = os.path.getsize(doc.file_path)
                await db.commit()

        result_lines = ["Excel modificado correctamente.", f"Celdas actualizadas: {applied}"]
        if errors:
            result_lines.append(f"Errores: {'; '.join(errors[:5])}")
        return "\n".join(result_lines)
    except Exception as e:
        return f"Error modificando Excel: {e}"


@tool
async def read_excel(
    tenant_id: str, document_id: str, sheet_name: str = "", max_rows: int = 30
) -> str:
    """
    Lee el contenido de un archivo Excel y lo devuelve en formato texto tabular.
    Útil para que el LLM vea los datos antes de decidir qué modificar.

    Args:
        tenant_id: ID del tenant
        document_id: ID del documento Excel
        sheet_name: Nombre de la hoja a leer (vacío = primera hoja)
        max_rows: Máximo de filas a devolver (por defecto 30)
    """
    return await _read_excel_async(tenant_id, document_id, sheet_name, max_rows)


async def _read_excel_async(
    tenant_id: str, document_id: str, sheet_name: str, max_rows: int
) -> str:
    try:
        doc, err = await _load_excel_doc(tenant_id, document_id)
        if err:
            return err

        wb = openpyxl.load_workbook(doc.file_path, read_only=True, data_only=True)

        if sheet_name:
            if sheet_name not in wb.sheetnames:
                wb.close()
                return f"Error: Hoja '{sheet_name}' no encontrada. Disponibles: {', '.join(wb.sheetnames)}"
            ws = wb[sheet_name]
        else:
            ws = wb.active

        rows = list(ws.iter_rows(values_only=True, max_row=max_rows + 1))
        wb.close()

        if not rows:
            return f"La hoja '{ws.title}' está vacía."

        # Formatear como tabla
        lines = [f"Hoja: {ws.title} | Hojas disponibles: {', '.join(wb.sheetnames)}"]
        for i, row in enumerate(rows):
            prefix = "HDR" if i == 0 else f"R{i}"
            cells = [str(c) if c is not None else "" for c in row]
            lines.append(f"  {prefix}: {' | '.join(cells)}")

        if len(rows) > max_rows:
            lines.append(f"  ... (mostrando {max_rows} de las filas disponibles)")

        return "\n".join(lines)
    except Exception as e:
        return f"Error leyendo Excel: {e}"


# ─── Lista de herramientas (exportada para agent.py) ─────────────────────────

tools = [
    export_erp_data,
    list_available_datasets,
    import_excel,
    modify_excel,
    read_excel,
    create_document,
    list_tenant_documents,
    get_document_content,
    get_tenant_knowledge,
    upsert_tenant_knowledge,
]
