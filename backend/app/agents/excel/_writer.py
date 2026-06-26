"""
Excel file writer utilities for the Excel agent.
Handles workbook creation, styling, and theme application.
"""

import os

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.security import sanitize_spreadsheet_cell


def _hex_to_lighter(hex_color: str, factor: float = 0.4) -> str:
    """Return a lighter version of a hex color by blending toward white."""
    h = hex_color.lstrip("#")
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    r = int(r + (255 - r) * factor)
    g = int(g + (255 - g) * factor)
    b = int(b + (255 - b) * factor)
    return f"{r:02X}{g:02X}{b:02X}"


def _write_excel(
    sheets: dict[str, pd.DataFrame], output_path: str, theme: dict | None = None
) -> None:
    _theme = theme or {}
    accent_hex = _theme.get("accent_color", "#1F4E79").lstrip("#")
    table_style = _theme.get("table_style", "striped")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(fill_type="solid", fgColor=accent_hex)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_side = Side(style="thin", color="D9D9D9")
    cell_border = Border(left=thin_side, right=thin_side, bottom=thin_side)

    if table_style == "minimal":
        alt_fill = None
    elif table_style == "bold":
        alt_fill = PatternFill(fill_type="solid", fgColor=_hex_to_lighter(accent_hex))
    else:  # "striped" (default)
        alt_fill = PatternFill(fill_type="solid", fgColor="EBF3FB")

    for sheet_name, df in sheets.items():
        ws = wb.create_sheet(title=sheet_name[:31])
        if df.empty:
            ws.append(["Sin datos"])
            continue
        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = cell_border
        for row_idx, row in enumerate(df.itertuples(index=False), start=2):
            fill = alt_fill if (alt_fill and row_idx % 2 == 0) else None
            for col_idx, value in enumerate(row, start=1):
                cell = ws.cell(
                    row=row_idx, column=col_idx, value=sanitize_spreadsheet_cell(value)
                )
                cell.border = cell_border
                cell.alignment = Alignment(horizontal="left", vertical="center")
                if fill:
                    cell.fill = fill
        for col_idx, col_name in enumerate(df.columns, start=1):
            max_len = max(
                len(str(col_name)),
                df.iloc[:, col_idx - 1].astype(str).str.len().max() if not df.empty else 0,
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 50)
        ws.freeze_panes = "A2"

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
