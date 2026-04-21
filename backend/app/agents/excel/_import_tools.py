"""
Excel import tools for the Excel agent.
Handles importing Excel data into the ERP database.
"""

import logging
import os
import uuid

import openpyxl
from langchain_core.tools import tool
from sqlalchemy.future import select

from app.db.base import AsyncSessionLocal
from app.db.models.crm import Client
from app.db.models.hr import Employee
from app.db.models.inventory import Product
from app.db.models.tenant import TenantDocument

logger = logging.getLogger(__name__)

# ─── Column mapping config ────────────────────────────────────────────────────

_IMPORT_COLUMN_MAP = {
    "clientes": {
        "model": "Client",
        "fields": {
            "nombre": "name",
            "name": "name",
            "nif": "nif",
            "cif": "nif",
            "email": "email",
            "correo": "email",
            "telefono": "phone",
            "teléfono": "phone",
            "phone": "phone",
            "direccion": "address",
            "dirección": "address",
            "address": "address",
        },
        "required": ["name"],
    },
    "productos": {
        "model": "Product",
        "fields": {
            "nombre": "name",
            "name": "name",
            "descripcion": "description",
            "descripción": "description",
            "description": "description",
            "precio": "price",
            "price": "price",
            "pvp": "price",
            "tipo": "item_type",
            "type": "item_type",
            "iva": "tax_percentage",
            "tax": "tax_percentage",
        },
        "required": ["name"],
    },
    "empleados": {
        "model": "Employee",
        "fields": {
            "nombre": "name",
            "name": "name",
            "nif": "nif",
            "dni": "nif",
            "email": "email",
            "correo": "email",
            "cargo": "role",
            "puesto": "role",
            "role": "role",
            "departamento": "department",
            "department": "department",
            "salario": "base_salary",
            "salario_base": "base_salary",
            "base_salary": "base_salary",
        },
        "required": ["name"],
    },
}


# ─── Private helpers ──────────────────────────────────────────────────────────


async def _load_excel_doc(
    tenant_id: str, document_id: str
) -> "tuple[TenantDocument | None, str | None]":
    """Carga TenantDocument desde BD y valida que el archivo exista en disco."""
    async with AsyncSessionLocal() as db:
        result = await db.execute(
            select(TenantDocument).where(
                TenantDocument.tenant_id == uuid.UUID(tenant_id),
                TenantDocument.id == uuid.UUID(document_id),
            )
        )
        doc = result.scalar_one_or_none()
    if not doc:
        return None, f"Error: Documento {document_id} no encontrado."
    if not doc.file_path or not os.path.exists(doc.file_path):
        return None, f"Error: Archivo no encontrado en disco: {doc.file_path}"
    return doc, None


def _detect_import_target(sheet_title: str) -> str | None:
    """Infiere el target de importacion a partir del nombre de la hoja."""
    title = sheet_title.lower().strip()
    for key in _IMPORT_COLUMN_MAP:
        if key in title:
            return key
    return None


def _map_excel_columns(headers: list[str], config: dict) -> dict[int, str]:
    """Mapea indices de columna Excel a nombres de campo del modelo."""
    return {idx: config["fields"][h] for idx, h in enumerate(headers) if h in config["fields"]}


async def _import_rows_to_db(
    rows: list, col_mapping: dict, config: dict, model_cls, tenant_id: str
) -> tuple[int, int, list[str]]:
    """Persiste filas Excel en la BD. Devuelve (created, skipped, errors)."""
    created, skipped, errors = 0, 0, []
    async with AsyncSessionLocal() as db:
        for row_idx, row in enumerate(rows[1:], start=2):
            record_data: dict = {"tenant_id": uuid.UUID(tenant_id)}
            for col_idx, field_name in col_mapping.items():
                value = row[col_idx] if col_idx < len(row) else None
                if value is not None:
                    record_data[field_name] = value
            if any(not record_data.get(r) for r in config["required"]):
                skipped += 1
                continue
            try:
                db.add(model_cls(**record_data))
                created += 1
            except Exception as e:
                errors.append(f"Fila {row_idx}: {e}")
                skipped += 1
        if created > 0:
            await db.commit()
    return created, skipped, errors


# ─── Tool ─────────────────────────────────────────────────────────────────────


@tool
async def import_excel(
    tenant_id: str, document_id: str, target: str = "", sheet_name: str = ""
) -> str:
    """
    Importa datos desde un archivo Excel (.xlsx) subido al sistema hacia la base de datos del ERP.
    Lee las columnas del Excel y las mapea automáticamente a campos de la BD.

    Args:
        tenant_id: ID del tenant
        document_id: ID del documento Excel a importar (usar list_tenant_documents para obtenerlo)
        target: Tipo de datos a importar ('clientes', 'productos', 'empleados'). Si vacío, se detecta por nombre de hoja.
        sheet_name: Nombre de la hoja a importar (vacío = primera hoja)
    """
    return await _import_excel_async(tenant_id, document_id, target, sheet_name)


async def _import_excel_async(
    tenant_id: str, document_id: str, target: str, sheet_name: str
) -> str:
    MODEL_MAP = {"Client": Client, "Product": Product, "Employee": Employee}

    try:
        doc, err = await _load_excel_doc(tenant_id, document_id)
        if err:
            return err

        wb = openpyxl.load_workbook(doc.file_path, read_only=True, data_only=True)

        if sheet_name:
            if sheet_name not in wb.sheetnames:
                return f"Error: Hoja '{sheet_name}' no encontrada. Hojas disponibles: {', '.join(wb.sheetnames)}"
            ws = wb[sheet_name]
        else:
            ws = wb.active

        target = (target or "").strip().lower() or _detect_import_target(ws.title) or ""
        if not target:
            return (
                f"Error: No se pudo detectar el tipo de datos de la hoja '{ws.title}'. "
                f"Especifica target: 'clientes', 'productos' o 'empleados'."
            )
        if target not in _IMPORT_COLUMN_MAP:
            return f"Error: Target '{target}' no soportado. Opciones: {', '.join(_IMPORT_COLUMN_MAP.keys())}"

        config = _IMPORT_COLUMN_MAP[target]
        model_cls = MODEL_MAP[config["model"]]

        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if len(rows) < 2:
            return "Error: El archivo no tiene datos (solo headers o vacío)."

        headers = [str(h).strip().lower() if h else "" for h in rows[0]]
        col_mapping = _map_excel_columns(headers, config)
        if not col_mapping:
            return (
                f"Error: No se reconocieron columnas para '{target}'. "
                f"Columnas encontradas: {', '.join(h for h in headers if h)}. "
                f"Columnas esperadas: {', '.join(config['fields'].keys())}"
            )

        sheet_title = ws.title
        created, skipped, errors = await _import_rows_to_db(
            rows, col_mapping, config, model_cls, tenant_id
        )

        result_lines = [
            f"Importación completada desde '{sheet_title}'.",
            f"Target: {target}",
            f"Creados: {created} registros",
            f"Omitidos: {skipped} filas (datos incompletos o errores)",
        ]
        if errors[:5]:
            result_lines.append(f"Errores: {'; '.join(errors[:5])}")
        return "\n".join(result_lines)
    except Exception as e:
        return f"Error importando Excel: {e}"
